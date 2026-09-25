#!/usr/bin/env python3
from __future__ import annotations
import hashlib, itertools, json, math, pathlib, urllib.request
from datetime import datetime, timezone

import numpy as np
import openpyxl

ROOT=pathlib.Path(__file__).resolve().parents[2]
BIO=ROOT/"BIO_CHI"
CFG_SRC=BIO/"config"/"SU2026_SUPPDATA_WORKBOOK_FREEZE_v0_1.json"
CFG_MAP=BIO/"config"/"SU2026_ONCOLOGY_SOURCE_MAPPING_FREEZE_v0_2.json"
CFG=BIO/"config"/"SU2026_NATIVE_COMPARATOR_SEM_LIMIT_FREEZE_v0_1.json"
OUTDIR=BIO/"artifacts"/"generated"
OUTDIR.mkdir(parents=True,exist_ok=True)
OUT=OUTDIR/"su2026_native_comparator_sem_limit_v0_1.json"
AUDIT=OUTDIR/"SU2026_NATIVE_COMPARATOR_SEM_LIMIT_V01_AUDIT.md"
UA="BioChiReviewerReproducibility/0.1"

def fetch(url):
    req=urllib.request.Request(url,headers={"User-Agent":UA,"Accept":"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,application/octet-stream,*/*;q=0.5"})
    with urllib.request.urlopen(req,timeout=180) as r:
        return r.read()

def md5(b): return hashlib.md5(b).hexdigest()
def sha256(b): return hashlib.sha256(b).hexdigest()

def acquire(cfg):
    errors=[]
    for url in cfg["candidate_urls"]:
        try:
            data=fetch(url)
            if md5(data)==cfg["expected_md5"]:
                return data
            errors.append({"url":url,"md5":md5(data),"bytes":len(data)})
        except Exception as e:
            errors.append({"url":url,"error":f"{type(e).__name__}: {e}"})
    raise RuntimeError(json.dumps(errors))

def as_num(v,coord):
    if isinstance(v,bool) or not isinstance(v,(int,float)): raise ValueError(f"{coord} not numeric: {v!r}")
    x=float(v)
    if not math.isfinite(x): raise ValueError(f"{coord} nonfinite")
    return x

def parse_sem(v,coord):
    if isinstance(v,(int,float)) and not isinstance(v,bool): x=float(v)
    elif isinstance(v,str): x=float(v.strip().replace("±","").replace("+/-","").strip())
    else: raise ValueError(f"{coord} invalid SEM {v!r}")
    if x<0 or not math.isfinite(x): raise ValueError(f"{coord} invalid SEM")
    return x

def A_from(p):
    return np.array([[p["Me-e"],p["Ml-e"]],[p["Me-l"],p["Ml-l"]]],dtype=float)

def metrics(A,thr=1e-12):
    vals,vecs=np.linalg.eig(A)
    tr=float(np.trace(A)); det=float(np.linalg.det(A)); disc=float(tr*tr-4*det)
    sab=float(np.max(np.real(vals)))
    cond=float(np.linalg.cond(vecs))
    complex_pair=bool(np.any(np.abs(np.imag(vals))>thr))
    chi=None
    if complex_pair:
        z=next(z for z in vals if abs(float(np.imag(z)))>thr)
        chi=float(-np.real(z)/abs(z))
    return {"trace":tr,"determinant":det,"discriminant":disc,"spectral_abscissa":sab,
            "eigenvector_condition_number":cond,"complex_pair_present":complex_pair,"chi":chi,
            "eigenvalues":[{"real":float(np.real(z)),"imag":float(np.imag(z))} for z in vals]}

def corner_summary(params,sem,level):
    dynamic=["Me-e","Ml-e","Me-l","Ml-l"]
    records=[]
    for signs in itertools.product([-1.0,1.0],repeat=4):
        p=dict(params)
        for k,s in zip(dynamic,signs):
            p[k]=params[k]+s*level*sem[k]
        m=metrics(A_from(p))
        records.append({"signs":dict(zip(dynamic,[int(s) for s in signs])),"params":{k:p[k] for k in dynamic},**m})
    d=[r["discriminant"] for r in records]
    sa=[r["spectral_abscissa"] for r in records]
    co=[r["eigenvector_condition_number"] for r in records]
    chis=[r["chi"] for r in records if r["chi"] is not None]
    return {
        "level_sem":level,
        "corner_count":len(records),
        "complex_corner_count":sum(r["complex_pair_present"] for r in records),
        "stable_corner_count":sum(r["spectral_abscissa"]<0 for r in records),
        "discriminant_min":min(d),"discriminant_max":max(d),
        "spectral_abscissa_min":min(sa),"spectral_abscissa_max":max(sa),
        "condition_number_min":min(co),"condition_number_max":max(co),
        "chi_min":min(chis) if chis else None,"chi_max":max(chis) if chis else None,
        "all_corners_complex":all(r["complex_pair_present"] for r in records),
        "no_corners_complex":not any(r["complex_pair_present"] for r in records),
        "all_corners_stable":all(r["spectral_abscissa"]<0 for r in records),
        "records":records
    }

def main():
    src=json.loads(CFG_SRC.read_text())
    mp=json.loads(CFG_MAP.read_text())
    cfg=json.loads(CFG.read_text())
    data=acquire(src)
    if sha256(data)!=mp["workbook_sha256"]: raise SystemExit("workbook SHA mismatch")
    path=OUTDIR/src["expected_filename"]; path.write_bytes(data)
    wb=openpyxl.load_workbook(path,read_only=False,data_only=False)
    ws=wb[mp["parameter_table"]["sheet"]]
    cols={"Bl":"D","Be":"E","Me-l":"F","Ml-l":"G","Ml-e":"H","Me-e":"I"}
    scenarios={}
    eq_diffs=[]
    for sc in mp["parameter_table"]["scenarios"]:
        row=sc["value_row"]; sr=sc["sem_row"]
        p={k:as_num(ws[f"{c}{row}"].value,f"{c}{row}") for k,c in cols.items()}
        se={k:parse_sem(ws[f"{c}{sr}"].value,f"{c}{sr}") for k,c in cols.items()}
        A=A_from(p); m=metrics(A)
        inv=None
        diff=None
        if m["complex_pair_present"] and m["determinant"]>0:
            inv=float(-m["trace"]/(2*math.sqrt(m["determinant"])))
            diff=float(abs(inv-m["chi"]))
            eq_diffs.append(diff)
        lim=[corner_summary(p,se,float(level)) for level in cfg["reported_sem_limit_map"]["levels"]]
        scenarios[sc["id"]]={"condition":sc["condition"],"mearly_polarity":sc["mearly_polarity"],"mlate_polarity":sc["mlate_polarity"],
                             "params":p,"sem":se,"nominal":m,"chi_from_invariants":inv,"chi_equivalence_abs_diff":diff,
                             "sem_limit_map":lim}
    wb.close()
    maxdiff=max(eq_diffs) if eq_diffs else None
    comparator_pass=(maxdiff is not None and maxdiff<=cfg["native_comparator"]["absolute_tolerance"])
    result={
      "schema_version":"0.1","generated_at_utc":datetime.now(timezone.utc).isoformat(),
      "gate":cfg["gate"],"status":"PASS" if comparator_pass else "FAIL_NATIVE_COMPARATOR_IDENTITY",
      "workbook_sha256":sha256(data),
      "native_comparator":{"status":"EXACT_EQUIVALENCE_CONFIRMED" if comparator_pass else "MISMATCH",
                           "max_abs_difference":maxdiff,
                           "identity":"chi = -trace(J)/(2*sqrt(det(J))) whenever the 2x2 real generator has a complex-conjugate pair"},
      "scenarios":scenarios,
      "interpretation":{
        "scalar_added_information_in_2D_source_model":"NONE_BEYOND_NORMALIZED_TRACE_AND_DETERMINANT",
        "scalar_role":"compact mode-regime coordinate equivalent to native pole geometry in this source model",
        "sem_limit_map_is_probability_interval":False
      }
    }
    OUT.write_text(json.dumps(result,indent=2)+"\n")
    lines=["# Su 2026 native-comparator and SEM Limit Map v0.1","",
           f"**Status:** {result['status']}",
           f"**Comparator:** {result['native_comparator']['status']}",
           f"**Maximum scalar identity error:** {maxdiff:.3e}" if maxdiff is not None else "**Maximum scalar identity error:** n/a","",
           "## Nominal and reported-SEM stress map","",
           "| Scenario | nominal complex | nominal chi | 1 SEM complex corners | 1 SEM stable | 2 SEM complex corners | 2 SEM stable |",
           "|---|---|---:|---:|---:|---:|---:|"]
    for sid,s in scenarios.items():
        chi="n/a" if s["nominal"]["chi"] is None else f"{s['nominal']['chi']:.8g}"
        l1=s["sem_limit_map"][0]; l2=s["sem_limit_map"][1]
        lines.append(f"| {sid} | {s['nominal']['complex_pair_present']} | {chi} | {l1['complex_corner_count']}/16 | {l1['stable_corner_count']}/16 | {l2['complex_corner_count']}/16 | {l2['stable_corner_count']}/16 |")
    lines += ["","## Interpretation","",
              "For a real 2x2 generator with a complex-conjugate pair, the scalar is algebraically equal to normalized trace/determinant pole geometry. It is therefore a compact coordinate, not additional information beyond those native invariants.",
              "",
              "The SEM grid is a deterministic hyperrectangle stress test, not a confidence interval and not a probability statement."]
    AUDIT.write_text("\n".join(lines)+"\n")
    print(json.dumps({"status":result["status"],"native_comparator":result["native_comparator"],
                      "summary":{sid:{"nominal_complex":s["nominal"]["complex_pair_present"],"nominal_chi":s["nominal"]["chi"],
                                      "sem1":{k:s["sem_limit_map"][0][k] for k in ["complex_corner_count","stable_corner_count","discriminant_min","discriminant_max","chi_min","chi_max"]},
                                      "sem2":{k:s["sem_limit_map"][1][k] for k in ["complex_corner_count","stable_corner_count","discriminant_min","discriminant_max","chi_min","chi_max"]}}
                                 for sid,s in scenarios.items()}},indent=2))
    if not comparator_pass: raise SystemExit("native comparator identity failed")
    print("BIO_CHI_SU2026_NATIVE_COMPARATOR_SEM_LIMIT_PASS")

if __name__=="__main__": main()
