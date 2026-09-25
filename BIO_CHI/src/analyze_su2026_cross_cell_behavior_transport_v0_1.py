#!/usr/bin/env python3
from __future__ import annotations
import hashlib, json, math, pathlib, urllib.request
from collections import Counter
from datetime import datetime, timezone

import numpy as np
import openpyxl
from scipy.integrate import solve_ivp

ROOT=pathlib.Path(__file__).resolve().parents[2]
BIO=ROOT/"BIO_CHI"
SRC=BIO/"config"/"SU2026_SUPPDATA_WORKBOOK_FREEZE_v0_1.json"
MAP=BIO/"config"/"SU2026_ONCOLOGY_SOURCE_MAPPING_FREEZE_v0_2.json"
CFG=BIO/"config"/"SU2026_CROSS_CELL_BEHAVIOR_TRANSPORT_FREEZE_v0_1.json"
OUTDIR=BIO/"artifacts"/"generated"
OUTDIR.mkdir(parents=True,exist_ok=True)
OUT=OUTDIR/"su2026_cross_cell_behavior_transport_v0_1.json"
AUDIT=OUTDIR/"SU2026_CROSS_CELL_BEHAVIOR_TRANSPORT_V01_AUDIT.md"
UA="BioChiReviewerReproducibility/0.1"

def fetch(url):
    req=urllib.request.Request(url,headers={"User-Agent":UA,"Accept":"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,application/octet-stream,*/*;q=0.5"})
    with urllib.request.urlopen(req,timeout=180) as r: return r.read()
def md5(b): return hashlib.md5(b).hexdigest()
def sha256(b): return hashlib.sha256(b).hexdigest()
def acquire(cfg):
    errs=[]
    for u in cfg["candidate_urls"]:
        try:
            b=fetch(u)
            if md5(b)==cfg["expected_md5"]: return b
            errs.append({"url":u,"md5":md5(b)})
        except Exception as e: errs.append({"url":u,"error":f"{type(e).__name__}: {e}"})
    raise RuntimeError(json.dumps(errs))
def num(v,coord):
    if isinstance(v,bool) or not isinstance(v,(int,float)): raise ValueError(f"{coord} not numeric: {v!r}")
    x=float(v)
    if not math.isfinite(x): raise ValueError(f"{coord} nonfinite")
    return x

def contribution_map(ws):
    out={}
    for r in range(3,ws.max_row+1):
        g=ws.cell(r,1).value
        if g is None: continue
        g=str(g).strip()
        if not g: continue
        if g in out: raise RuntimeError(f"duplicate contribution gene {g}")
        out[g]=[num(ws.cell(r,2).value,ws.cell(r,2).coordinate),num(ws.cell(r,3).value,ws.cell(r,3).coordinate)]
    return out
def select(mp,idx,pol):
    vals=[]
    for g,scores in mp.items():
        s=scores[idx]
        if pol=="positive" and s>0: vals.append((g,s))
        elif pol=="negative" and s<0: vals.append((g,s))
    vals.sort(key=(lambda x:(-x[1],x[0])) if pol=="positive" else (lambda x:(x[1],x[0])))
    genes=[g for g,_ in vals[:500]]
    if len(genes)!=500: raise RuntimeError(f"only {len(genes)} genes selected")
    return genes

def sheet_gene_rows(ws):
    rows={}
    for r in range(3,ws.max_row+1):
        v=ws.cell(r,1).value
        if v is None: continue
        g=str(v).strip()
        if g: rows.setdefault(g,[]).append(r)
    return rows

def unique_expression_map(ws, gene_rows, labels_to_cols, allowed_genes):
    out={}
    for g in allowed_genes:
        rs=gene_rows.get(g,[])
        if len(rs)!=1: raise RuntimeError(f"{ws.title}: expected one row for {g}, got {len(rs)}")
        r=rs[0]
        out[g]={}
        for label,col_letter in labels_to_cols.items():
            c=openpyxl.utils.column_index_from_string(col_letter)
            out[g][label]=num(ws.cell(r,c).value,ws.cell(r,c).coordinate)
    return out

def mean_state(expr,egenes,lgenes,label):
    e=np.array([expr[g][label] for g in egenes],dtype=float)
    l=np.array([expr[g][label] for g in lgenes],dtype=float)
    return np.array([float(e.mean()),float(l.mean())],dtype=float)

def generator_from_row(ws,row):
    vals={k:num(ws[f"{c}{row}"].value,f"{c}{row}") for k,c in {"Bl":"D","Be":"E","Me-l":"F","Ml-l":"G","Ml-e":"H","Me-e":"I"}.items()}
    A=np.array([[vals["Me-e"],vals["Ml-e"]],[vals["Me-l"],vals["Ml-l"]]],dtype=float)
    b=np.array([vals["Be"],vals["Bl"]],dtype=float)
    return vals,A,b

def simulate(A,b,times,x0):
    def rhs(t,x): return b+A@x
    sol=solve_ivp(rhs,(float(times[0]),float(times[-1])),x0,t_eval=np.array(times,dtype=float),method="RK45",max_step=0.1,rtol=1e-10,atol=1e-12)
    if not sol.success or sol.y.shape[1]!=len(times): raise RuntimeError(sol.message)
    y=sol.y.T
    if not np.all(np.isfinite(y)): raise RuntimeError("nonfinite prediction")
    return y

def norm_sse(obs,pred,skip0=True):
    o=obs[1:] if skip0 else obs
    p=pred[1:] if skip0 else pred
    total=0.0; states=[]
    for k,name in enumerate(["Mearly","Mlate"]):
        y=o[:,k]; z=p[:,k]
        sse=float(np.sum((z-y)**2))
        denom=float(np.sum((y-np.mean(y))**2))
        nsse=sse/denom if denom>0 else (0.0 if sse==0 else float("inf"))
        rmse=float(np.sqrt(np.mean((z-y)**2)))
        corr=float(np.corrcoef(y,z)[0,1]) if np.std(y)>0 and np.std(z)>0 else None
        total+=nsse
        states.append({"state":name,"sse":sse,"variance_normalized_sse":nsse,"rmse":rmse,"pearson":corr})
    return {"aggregate_statewise_variance_normalized_sse":total,"states":states}

def cosine(a,b):
    na=float(np.linalg.norm(a)); nb=float(np.linalg.norm(b))
    if na==0 or nb==0: return None
    return float(np.dot(a,b)/(na*nb))

def main():
    src=json.loads(SRC.read_text())
    mp=json.loads(MAP.read_text())
    cfg=json.loads(CFG.read_text())
    data=acquire(src)
    if sha256(data)!=mp["workbook_sha256"]: raise SystemExit("workbook SHA mismatch")
    path=OUTDIR/src["expected_filename"]; path.write_bytes(data)
    wb=openpyxl.load_workbook(path,read_only=False,data_only=False)

    cm=contribution_map(wb["Supplementary Data 4"])
    gene_sets={
      "Mearly_positive":select(cm,1,"positive"),"Mearly_negative":select(cm,1,"negative"),
      "Mlate_positive":select(cm,0,"positive"),"Mlate_negative":select(cm,0,"negative")
    }

    # Source M397 rows are unique under the already executed source analysis.
    ws_m=wb["Supplementary Data 1"]
    mrows=sheet_gene_rows(ws_m)
    source_cols={"D0":"C","D3":"D","D8":"E","D13":"F","D21":"G","D29":"H","D33":"I","D38":"J","D59":"K"}
    scenarios_cfg={x["id"]:x for x in mp["parameter_table"]["scenarios"] if x["condition"]=="Drug ON"}
    ws5=wb["Supplementary Data 5"]

    results={}
    aggregate_all={}
    for cand in cfg["candidate_systems_in_order"]:
        ws=wb[cand["sheet"]]
        rows=sheet_gene_rows(ws)
        duplicate_symbols={g for g,rs in rows.items() if len(rs)>1}
        coverage={}
        for name,genes in gene_sets.items():
            unamb=[g for g in genes if g in rows and len(rows[g])==1]
            coverage[name]={
              "source_n":500,"unambiguous_present_n":len(unamb),"coverage_fraction":len(unamb)/500.0,
              "ambiguous_selected_n":sum(1 for g in genes if g in duplicate_symbols),
              "missing_selected_n":sum(1 for g in genes if g not in rows),
              "unambiguous_genes":unamb
            }
        eligible=all(v["coverage_fraction"]>=cfg["source_eligibility"]["minimum_unambiguous_coverage_each_of_four_source_gene_sets"] for v in coverage.values())
        sysrec={"eligible":eligible,"coverage":{k:{kk:vv for kk,vv in v.items() if kk!="unambiguous_genes"} for k,v in coverage.items()},
                "expression_values_read":False}
        if not eligible:
            sysrec["disposition"]="SOURCE_LIMITED_NOT_EVALUATED_VALUES_REMAIN_UNOPENED"
            results[cand["system"]]=sysrec
            continue

        # Only now are target expression values read.
        needed=set()
        for v in coverage.values(): needed.update(v["unambiguous_genes"])
        tcols={lab:openpyxl.utils.get_column_letter(i+2) for i,lab in enumerate(cand["labels"])}
        # Explicit columns B:E correspond Control,D3,D14,D21 under frozen schema.
        tcols=dict(zip(cand["labels"],["B","C","D","E"]))
        expr_t=unique_expression_map(ws,rows,tcols,needed)
        expr_m=unique_expression_map(ws_m,mrows,source_cols,needed)
        sysrec["expression_values_read"]=True
        carriers={}
        aggregate={"generator":0.0,"constant":0.0,"empirical_template":0.0}
        for tag,pols in {"PP":("positive","positive"),"NP":("negative","positive"),"PN":("positive","negative"),"NN":("negative","negative")}.items():
            epol,lpol=pols
            egenes=coverage[f"Mearly_{epol}"]["unambiguous_genes"]
            lgenes=coverage[f"Mlate_{lpol}"]["unambiguous_genes"]
            labels=cand["labels"]; times=cand["times_days"]
            obs=np.vstack([mean_state(expr_t,egenes,lgenes,lab) for lab in labels])
            srcstates={lab:mean_state(expr_m,egenes,lgenes,lab) for lab in source_cols}
            _,A,b=generator_from_row(ws5,scenarios_cfg[f"ON_{tag}"]["value_row"])
            gen=simulate(A,b,times,obs[0])
            const=np.repeat(obs[[0]],len(times),axis=0)
            source_map={"D3":"D3","D14":"D13","D21":"D21"}
            templ=[obs[0]]
            for lab in labels[1:]:
                sl=source_map[lab]
                templ.append(obs[0]+(srcstates[sl]-srcstates["D0"]))
            templ=np.vstack(templ)
            mg=norm_sse(obs,gen); mc=norm_sse(obs,const); mt=norm_sse(obs,templ)
            aggregate["generator"]+=mg["aggregate_statewise_variance_normalized_sse"]
            aggregate["constant"]+=mc["aggregate_statewise_variance_normalized_sse"]
            aggregate["empirical_template"]+=mt["aggregate_statewise_variance_normalized_sse"]
            carriers[tag]={
              "n_genes":{"Mearly":len(egenes),"Mlate":len(lgenes)},
              "observed_states":obs.tolist(),"generator_prediction":gen.tolist(),
              "constant_prediction":const.tolist(),"empirical_template_prediction":templ.tolist(),
              "metrics":{"generator":mg,"constant":mc,"empirical_template":mt},
              "D3_cosine_generator":cosine(gen[1]-gen[0],obs[1]-obs[0]),
              "D3_cosine_empirical_template":cosine(templ[1]-templ[0],obs[1]-obs[0])
            }
        passed=aggregate["generator"]<aggregate["constant"] and aggregate["generator"]<aggregate["empirical_template"]
        sysrec["aggregate_metric"]=aggregate
        sysrec["primary_gate_pass"]=passed
        sysrec["carriers"]=carriers
        sysrec["disposition"]="PASS_NO_RETUNING_BEHAVIOR_TRANSPORT" if passed else "FAIL_NO_RETUNING_BEHAVIOR_TRANSPORT"
        results[cand["system"]]=sysrec
        aggregate_all[cand["system"]]=aggregate

    wb.close()
    eligible_systems=[k for k,v in results.items() if v["eligible"]]
    pass_systems=[k for k,v in results.items() if v.get("primary_gate_pass") is True]
    overall="PASS_C3_HELDOUT_BEHAVIOR_GATE" if pass_systems else ("FAIL_C3_HELDOUT_BEHAVIOR_GATE" if eligible_systems else "NO_ELIGIBLE_CROSS_CELL_SYSTEM")
    result={
      "schema_version":"0.1","generated_at_utc":datetime.now(timezone.utc).isoformat(),
      "gate":cfg["gate"],"status":overall,"workbook_sha256":sha256(data),
      "eligible_systems":eligible_systems,"passing_systems":pass_systems,"systems":results,
      "epistemic":{"tier":"P0-Q","P1_confirmation":False,"target_scalar_constructed":False,
                   "no_parameter_refitting":True,"all_four_carriers_retained":True}
    }
    OUT.write_text(json.dumps(result,indent=2)+"\n")
    lines=["# Su 2026 cross-cell behavior transport v0.1","",f"**Status:** {overall}","",
           "| System | Eligible | Values read | Generator aggregate | Constant aggregate | Empirical-template aggregate | Gate |",
           "|---|---|---|---:|---:|---:|---|"]
    for s,v in results.items():
        if v["eligible"]:
            a=v["aggregate_metric"]
            lines.append(f"| {s} | yes | yes | {a['generator']:.8g} | {a['constant']:.8g} | {a['empirical_template']:.8g} | {v['disposition']} |")
        else:
            lines.append(f"| {s} | no | no | n/a | n/a | n/a | {v['disposition']} |")
    lines += ["","No target-system scalar was constructed. This is a no-retuning behavior test of the frozen M397 drug-on generator family."]
    AUDIT.write_text("\n".join(lines)+"\n")
    print(json.dumps({
      "status":overall,
      "systems":{k:{"eligible":v["eligible"],"expression_values_read":v["expression_values_read"],
                    "coverage":v["coverage"],"aggregate_metric":v.get("aggregate_metric"),
                    "primary_gate_pass":v.get("primary_gate_pass"),"disposition":v["disposition"]} for k,v in results.items()}
    },indent=2))
    print("BIO_CHI_SU2026_CROSS_CELL_BEHAVIOR_TRANSPORT_DONE")

if __name__=="__main__": main()
