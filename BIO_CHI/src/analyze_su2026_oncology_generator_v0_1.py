#!/usr/bin/env python3
from __future__ import annotations

import hashlib
import json
import math
import pathlib
import urllib.request
from datetime import datetime, timezone

import numpy as np
import openpyxl
from scipy.integrate import solve_ivp
from scipy.linalg import expm

ROOT = pathlib.Path(__file__).resolve().parents[2]
BIO = ROOT / "BIO_CHI"
CFG_SOURCE = BIO / "config" / "SU2026_SUPPDATA_WORKBOOK_FREEZE_v0_1.json"
CFG_MAP = BIO / "config" / "SU2026_ONCOLOGY_SOURCE_MAPPING_FREEZE_v0_2.json"
CFG_ANALYSIS = BIO / "config" / "SU2026_ONCOLOGY_GENERATOR_ANALYSIS_FREEZE_v0_1.json"
OUTDIR = BIO / "artifacts" / "generated"
OUTDIR.mkdir(parents=True, exist_ok=True)
OUT_JSON = OUTDIR / "su2026_oncology_generator_analysis_v0_1.json"
OUT_MD = OUTDIR / "SU2026_ONCOLOGY_GENERATOR_ANALYSIS_V01_AUDIT.md"
UA = "BioChiReviewerReproducibility/0.1"

def fetch(url: str) -> bytes:
    req = urllib.request.Request(url, headers={
        "User-Agent": UA,
        "Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,application/octet-stream,*/*;q=0.5",
    })
    with urllib.request.urlopen(req, timeout=180) as r:
        return r.read()

def md5(b: bytes) -> str:
    return hashlib.md5(b).hexdigest()

def sha256(b: bytes) -> str:
    return hashlib.sha256(b).hexdigest()

def download_verified(cfg: dict) -> bytes:
    errors = []
    for url in cfg["candidate_urls"]:
        try:
            data = fetch(url)
            if md5(data) != cfg["expected_md5"]:
                errors.append({"url": url, "error": "MD5_MISMATCH", "md5": md5(data), "bytes": len(data)})
                continue
            return data
        except Exception as exc:
            errors.append({"url": url, "error": f"{type(exc).__name__}: {exc}"})
    raise RuntimeError("workbook acquisition failed: " + json.dumps(errors))

def num(v, coord: str) -> float:
    if isinstance(v, bool) or not isinstance(v, (int, float)):
        raise ValueError(f"expected numeric cell at {coord}, got {v!r}")
    if not math.isfinite(float(v)):
        raise ValueError(f"nonfinite numeric at {coord}")
    return float(v)

def parse_sem(v, coord: str) -> float:
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        x = float(v)
    elif isinstance(v, str):
        s = v.strip().replace("±", "").replace("+/-", "").strip()
        x = float(s)
    else:
        raise ValueError(f"unparseable SEM at {coord}: {v!r}")
    if not math.isfinite(x) or x < 0:
        raise ValueError(f"invalid SEM at {coord}: {x}")
    return x

def build_gene_map(ws, score_cols):
    out = {}
    duplicates = []
    for r in range(3, ws.max_row + 1):
        gene = ws.cell(r, 1).value
        if gene is None:
            continue
        gene = str(gene).strip()
        if not gene:
            continue
        if gene in out:
            duplicates.append(gene)
            continue
        vals = []
        for c in score_cols:
            vals.append(num(ws.cell(r, c).value, ws.cell(r, c).coordinate))
        out[gene] = vals
    return out, sorted(set(duplicates))

def select_polarity(score_map, idx, polarity):
    rows = []
    for gene, vals in score_map.items():
        s = vals[idx]
        if polarity == "positive" and s > 0:
            rows.append((gene, s))
        elif polarity == "negative" and s < 0:
            rows.append((gene, s))
    if polarity == "positive":
        rows.sort(key=lambda x: (-x[1], x[0]))
    else:
        rows.sort(key=lambda x: (x[1], x[0]))
    return [g for g, _ in rows[:500]], rows

def build_expression_map(ws, columns):
    out = {}
    duplicates = []
    for r in range(3, ws.max_row + 1):
        gene = ws.cell(r, 1).value
        if gene is None:
            continue
        gene = str(gene).strip()
        if not gene:
            continue
        if gene in out:
            duplicates.append(gene)
            continue
        vals = {}
        for label, col_letter in columns.items():
            c = openpyxl.utils.column_index_from_string(col_letter)
            vals[label] = num(ws.cell(r, c).value, ws.cell(r, c).coordinate)
        out[gene] = vals
    return out, sorted(set(duplicates))

def mean_state(expr, genes_early, genes_late, label):
    missing_e = [g for g in genes_early if g not in expr]
    missing_l = [g for g in genes_late if g not in expr]
    if missing_e or missing_l:
        raise ValueError(f"missing selected genes in expression sheet at {label}: early={len(missing_e)}, late={len(missing_l)}")
    e = np.array([expr[g][label] for g in genes_early], dtype=float)
    l = np.array([expr[g][label] for g in genes_late], dtype=float)
    return np.array([float(np.mean(e)), float(np.mean(l))], dtype=float)

def generator(params, mapping):
    Bl=params["Bl"]; Be=params["Be"]; Mel=params["Me-l"]; Mll=params["Ml-l"]; Mle=params["Ml-e"]; Mee=params["Me-e"]
    b=np.array([Be, Bl], dtype=float)
    if mapping == "PUBLICATION_EQUATION":
        A=np.array([[Mee, Mle],[Mel, Mll]], dtype=float)
    elif mapping == "PUBLIC_CODE_MAPPING":
        A=np.array([[-Mee, Mle],[-Mel, -Mll]], dtype=float)
    else:
        raise KeyError(mapping)
    return A,b

def simulate(A,b,times,x0):
    def rhs(t,x):
        return b + A @ x
    sol=solve_ivp(rhs,(float(times[0]),float(times[-1])),x0,t_eval=np.array(times,dtype=float),method="RK45",max_step=0.1,rtol=1e-10,atol=1e-12)
    if not sol.success or sol.y.shape[1] != len(times):
        raise RuntimeError("RK45 failure: "+sol.message)
    pred=sol.y.T
    # exact affine crosscheck via augmented exponential
    aug=np.zeros((3,3),dtype=float)
    aug[:2,:2]=A
    aug[:2,2]=b
    exact=[]
    z0=np.array([x0[0],x0[1],1.0])
    for t in times:
        exact.append((expm(aug*float(t))@z0)[:2])
    exact=np.array(exact)
    maxdiff=float(np.max(np.abs(pred-exact)))
    scale=max(1.0,float(np.max(np.abs(exact))))
    rel=maxdiff/scale
    if rel > 1e-7:
        raise RuntimeError(f"RK45/closed-form mismatch rel={rel}")
    return pred, rel

def fit_metrics(obs,pred):
    resid=pred-obs
    rmse=float(np.sqrt(np.mean(resid**2)))
    state=[]
    norm_sse=0.0
    for k,name in enumerate(["Mearly","Mlate"]):
        y=obs[:,k]; p=pred[:,k]
        r=p-y
        sse=float(np.sum(r*r))
        denom=float(np.sum((y-np.mean(y))**2))
        norm_sse += sse/denom if denom>0 else (0.0 if sse==0 else float("inf"))
        rng=float(np.max(y)-np.min(y))
        nrmse=float(np.sqrt(np.mean(r*r))/rng) if rng>0 else (0.0 if np.allclose(r,0) else float("inf"))
        pear=float(np.corrcoef(y,p)[0,1]) if np.std(y)>0 and np.std(p)>0 else None
        # simple Spearman via ranks
        ry=np.argsort(np.argsort(y)).astype(float)
        rp=np.argsort(np.argsort(p)).astype(float)
        spear=float(np.corrcoef(ry,rp)[0,1]) if np.std(ry)>0 and np.std(rp)>0 else None
        state.append({"state":name,"rmse":float(np.sqrt(np.mean(r*r))),"nrmse_range":nrmse,"pearson":pear,"spearman":spear,"sse":sse,"variance_normalized_sse":(sse/denom if denom>0 else None)})
    return {"joint_rmse":rmse,"aggregate_statewise_variance_normalized_sse":norm_sse,"states":state}

def modal_metrics(A,b,complex_thr=1e-10):
    vals, vecs=np.linalg.eig(A)
    order=np.argsort(np.real(vals))[::-1]
    vals=vals[order]; vecs=vecs[:,order]
    tr=float(np.trace(A)); det=float(np.linalg.det(A)); disc=float(tr*tr-4*det)
    cond=float(np.linalg.cond(vecs))
    sab=float(np.max(np.real(vals)))
    complex_present=bool(np.any(np.abs(np.imag(vals))>complex_thr))
    chi=None
    pair=None
    if complex_present:
        cps=[z for z in vals if abs(float(np.imag(z)))>complex_thr]
        if len(cps)!=2 or not np.allclose(cps[0],np.conj(cps[1]),rtol=1e-7,atol=1e-9):
            raise RuntimeError("non-conjugate complex spectrum in real 2x2 generator")
        z=cps[0]
        chi=float(-np.real(z)/abs(z))
        pair=[{"real":float(np.real(v)),"imag":float(np.imag(v))} for v in cps]
    equilibrium=None
    if abs(det)>1e-12:
        equilibrium=(-np.linalg.solve(A,b)).tolist()
    return {
        "A":A.tolist(),"b":b.tolist(),"trace":tr,"determinant":det,"discriminant":disc,
        "eigenvalues":[{"real":float(np.real(v)),"imag":float(np.imag(v))} for v in vals],
        "eigenvector_condition_number":cond,"spectral_abscissa":sab,
        "complex_pair_present":complex_present,"complex_pair":pair,"chi_bio_candidate":chi,
        "equilibrium":equilibrium
    }

def main():
    src=json.loads(CFG_SOURCE.read_text())
    mapping_cfg=json.loads(CFG_MAP.read_text())
    analysis_cfg=json.loads(CFG_ANALYSIS.read_text())
    data=download_verified(src)
    observed_sha=sha256(data)
    if observed_sha != mapping_cfg["workbook_sha256"]:
        raise SystemExit(f"SHA256 mismatch: {observed_sha}")
    local=OUTDIR/src["expected_filename"]
    local.write_bytes(data)

    wb=openpyxl.load_workbook(local,read_only=False,data_only=False)
    ws5=wb["Supplementary Data 5"]
    ws4=wb["Supplementary Data 4"]
    ws1=wb["Supplementary Data 1"]

    # Validate frozen headers.
    for name,coord in mapping_cfg["parameter_table"]["headers"].items():
        if str(ws5[coord].value).strip()!=name:
            raise SystemExit(f"header mismatch {coord}: {ws5[coord].value!r} != {name!r}")

    score_map,dup4=build_gene_map(ws4,[2,3])
    expr_map,dup1=build_expression_map(ws1,mapping_cfg["state_reconstruction"]["expression_time_columns"])
    if dup4:
        raise SystemExit(f"duplicate contribution genes: {dup4[:20]}")
    if dup1:
        raise SystemExit(f"duplicate expression genes: {dup1[:20]}")

    gene_sets={}
    for module,idx in [("Mlate",0),("Mearly",1)]:
        for pol in ["positive","negative"]:
            genes,eligible=select_polarity(score_map,idx,pol)
            if len(genes)!=500:
                raise SystemExit(f"{module} {pol}: only {len(genes)} selected")
            missing=[g for g in genes if g not in expr_map]
            if missing:
                raise SystemExit(f"{module} {pol}: selected genes missing from expression: {missing[:20]}")
            gene_sets[f"{module}_{pol}"]={
                "n":len(genes),
                "genes_sha256":hashlib.sha256("\n".join(genes).encode()).hexdigest(),
                "score_min":float(min(score_map[g][idx] for g in genes)),
                "score_max":float(max(score_map[g][idx] for g in genes))
            }

    scenarios={}
    comparisons={"PUBLICATION_EQUATION":[],"PUBLIC_CODE_MAPPING":[]}
    aggregate={"PUBLICATION_EQUATION":0.0,"PUBLIC_CODE_MAPPING":0.0}
    for sc in mapping_cfg["parameter_table"]["scenarios"]:
        rid=sc["id"]; row=sc["value_row"]; semrow=sc["sem_row"]
        cols={"Bl":"D","Be":"E","Me-l":"F","Ml-l":"G","Ml-e":"H","Me-e":"I"}
        params={k:num(ws5[f"{col}{row}"].value,f"{col}{row}") for k,col in cols.items()}
        sem={k:parse_sem(ws5[f"{col}{semrow}"].value,f"{col}{semrow}") for k,col in cols.items()}
        epol=sc["mearly_polarity"]; lpol=sc["mlate_polarity"]
        egenes=select_polarity(score_map,1,epol)[0]
        lgenes=select_polarity(score_map,0,lpol)[0]
        proto=mapping_cfg["trajectory_protocol"][sc["condition"]]
        labels=proto["labels"]; times=proto["times_days"]
        obs=np.vstack([mean_state(expr_map,egenes,lgenes,label) for label in labels])
        x0=obs[0].copy()
        maps={}
        for mname in ["PUBLICATION_EQUATION","PUBLIC_CODE_MAPPING"]:
            A,b=generator(params,mname)
            pred,cross=simulate(A,b,times,x0)
            fm=fit_metrics(obs,pred)
            mm=modal_metrics(A,b,analysis_cfg["modal_checks"]["complex_threshold_abs_imag"])
            maps[mname]={
                "fit":fm,"modal":mm,"rk45_closed_form_relative_max_error":cross,
                "predicted_states":pred.tolist()
            }
            comparisons[mname].append(fm["joint_rmse"])
            aggregate[mname]+=fm["aggregate_statewise_variance_normalized_sse"]
        scenarios[rid]={
            "condition":sc["condition"],"mearly_polarity":epol,"mlate_polarity":lpol,
            "params":params,"sem":sem,"times_days":times,"labels":labels,
            "observed_states":obs.tolist(),"mappings":maps
        }

    pub_better=all(scenarios[r]["mappings"]["PUBLICATION_EQUATION"]["fit"]["joint_rmse"] < scenarios[r]["mappings"]["PUBLIC_CODE_MAPPING"]["fit"]["joint_rmse"] for r in scenarios)
    code_better=all(scenarios[r]["mappings"]["PUBLIC_CODE_MAPPING"]["fit"]["joint_rmse"] < scenarios[r]["mappings"]["PUBLICATION_EQUATION"]["fit"]["joint_rmse"] for r in scenarios)
    if pub_better and aggregate["PUBLICATION_EQUATION"] < aggregate["PUBLIC_CODE_MAPPING"]:
        selected="PUBLICATION_EQUATION"
        mapping_status="UNIQUE_MAPPING_SOURCE_FIDELITY_SUPPORTED"
    elif code_better and aggregate["PUBLIC_CODE_MAPPING"] < aggregate["PUBLICATION_EQUATION"]:
        selected="PUBLIC_CODE_MAPPING"
        mapping_status="UNIQUE_MAPPING_SOURCE_FIDELITY_SUPPORTED"
    else:
        selected=None
        mapping_status="REFUSE_UNIQUE_GENERATOR_NONIDENTIFIABLE"

    cross={}
    for tag in ["PP","NP","PN","NN"]:
        on=scenarios[f"ON_{tag}"]; off=scenarios[f"OFF_{tag}"]
        rec={"selected_mapping":selected}
        if selected:
            mon=on["mappings"][selected]["modal"]; mof=off["mappings"][selected]["modal"]
            Aon=np.array(mon["A"]); Aoff=np.array(mof["A"])
            denom=np.linalg.norm(Aon,"fro")
            rec["relative_Frobenius_generator_change"]=float(np.linalg.norm(Aoff-Aon,"fro")/denom) if denom>0 else None
            rec["trace_change"]=float(mof["trace"]-mon["trace"])
            rec["determinant_change"]=float(mof["determinant"]-mon["determinant"])
            rec["discriminant_change"]=float(mof["discriminant"]-mon["discriminant"])
            if mon["chi_bio_candidate"] is not None and mof["chi_bio_candidate"] is not None:
                rec["scalar_motion_testable"]=True
                rec["chi_on"]=mon["chi_bio_candidate"]
                rec["chi_off"]=mof["chi_bio_candidate"]
                rec["delta_chi_off_minus_on"]=float(mof["chi_bio_candidate"]-mon["chi_bio_candidate"])
            else:
                rec["scalar_motion_testable"]=False
                rec["chi_on"]=mon["chi_bio_candidate"]
                rec["chi_off"]=mof["chi_bio_candidate"]
                rec["delta_chi_off_minus_on"]=None
            rec["memory_disposition"]="NOT_TESTABLE_FROM_SEPARATE_DRUG_ON_AND_DRUG_OFF_GENERATORS_ALONE"
        cross[tag]=rec

    result={
        "schema_version":"0.1",
        "generated_at_utc":datetime.now(timezone.utc).isoformat(),
        "gate":analysis_cfg["gate"],
        "status":"PASS_ANALYSIS_EXECUTED" if selected else "SOURCE_GENERATOR_MAPPING_REFUSED",
        "source_workbook":{"md5":md5(data),"sha256":observed_sha,"bytes":len(data)},
        "gene_sets":gene_sets,
        "mapping_adjudication":{
            "status":mapping_status,"selected_mapping":selected,
            "aggregate_statewise_variance_normalized_sse":aggregate,
            "scenario_joint_rmse":{
                rid:{m:scenarios[rid]["mappings"][m]["fit"]["joint_rmse"] for m in ["PUBLICATION_EQUATION","PUBLIC_CODE_MAPPING"]}
                for rid in scenarios
            }
        },
        "scenarios":scenarios,
        "cross_condition":cross,
        "epistemic":{
            "tier":"P0-Q",
            "P1_confirmation":False,
            "chi_equal_1_boundary":False,
            "memory_claim":"NOT_TESTABLE_FROM_SEPARATE_GENERATORS_ALONE",
            "existing_M397_recovery_result":"P0-D post-result independent evidence layer"
        }
    }
    if selected:
        complex_count=sum(1 for s in scenarios.values() if s["mappings"][selected]["modal"]["complex_pair_present"])
        result["qualification_summary"]={
            "selected_generator_mapping":selected,
            "scenario_count":8,
            "complex_scenario_count":complex_count,
            "real_only_scalar_refusal_count":8-complex_count,
            "scalar_representation_present_in_oncology_model":complex_count>0,
            "complete_two_state_modal_records_finite":all(np.isfinite(s["mappings"][selected]["modal"]["eigenvector_condition_number"]) for s in scenarios.values())
        }
    else:
        result["qualification_summary"]={
            "selected_generator_mapping":None,
            "scalar_admission":"BLOCKED_BY_SOURCE_GENERATOR_COLLISION",
            "modal_admission":"BLOCKED_BY_SOURCE_GENERATOR_COLLISION"
        }

    wb.close()
    OUT_JSON.write_text(json.dumps(result,indent=2)+"\n",encoding="utf-8")

    lines=[
        "# Su 2026 oncology generator analysis v0.1",
        "",
        f"**Status:** {result['status']}",
        f"**Mapping adjudication:** {mapping_status}",
        f"**Selected mapping:** {selected or 'NONE'}",
        "",
        "## Mapping fit comparison",
        "",
        "| Scenario | Publication-equation RMSE | Public-code RMSE |",
        "|---|---:|---:|",
    ]
    for rid in scenarios:
        a=scenarios[rid]["mappings"]["PUBLICATION_EQUATION"]["fit"]["joint_rmse"]
        b=scenarios[rid]["mappings"]["PUBLIC_CODE_MAPPING"]["fit"]["joint_rmse"]
        lines.append(f"| {rid} | {a:.8g} | {b:.8g} |")
    lines += ["", "Aggregate statewise variance-normalized SSE:",
              f"- publication equation: {aggregate['PUBLICATION_EQUATION']:.8g}",
              f"- public code mapping: {aggregate['PUBLIC_CODE_MAPPING']:.8g}", ""]
    if selected:
        lines += ["## Selected-mapping modal inventory","",
                  "| Scenario | trace | determinant | discriminant | spectral abscissa | complex pair | chi candidate |",
                  "|---|---:|---:|---:|---:|---|---:|"]
        for rid,s in scenarios.items():
            m=s["mappings"][selected]["modal"]
            chi="n/a" if m["chi_bio_candidate"] is None else f"{m['chi_bio_candidate']:.8g}"
            lines.append(f"| {rid} | {m['trace']:.8g} | {m['determinant']:.8g} | {m['discriminant']:.8g} | {m['spectral_abscissa']:.8g} | {m['complex_pair_present']} | {chi} |")
        lines += ["","## Frozen cross-condition comparisons","",
                  "| Carrier | relative generator change | scalar motion testable | delta chi (OFF-ON) |",
                  "|---|---:|---|---:|"]
        for tag,rec in cross.items():
            dc="n/a" if rec.get("delta_chi_off_minus_on") is None else f"{rec['delta_chi_off_minus_on']:.8g}"
            rg=rec.get("relative_Frobenius_generator_change")
            lines.append(f"| {tag} | {rg:.8g} | {rec.get('scalar_motion_testable')} | {dc} |")
    lines += ["","## Ceiling","","P0-Q only. No universal biological scalar value or chi=1 boundary is inferred. Separate drug-on and drug-off fitted generators do not by themselves license a memory claim."]
    OUT_MD.write_text("\n".join(lines)+"\n",encoding="utf-8")

    print(json.dumps({
        "status":result["status"],
        "mapping":result["mapping_adjudication"],
        "qualification_summary":result["qualification_summary"],
        "cross_condition":result["cross_condition"]
    },indent=2))
    print("BIO_CHI_SU2026_ONCOLOGY_GENERATOR_ANALYSIS_DONE")

if __name__=="__main__":
    main()
