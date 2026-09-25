#!/usr/bin/env python3
from __future__ import annotations
import hashlib, json, math, pathlib, urllib.request
from datetime import datetime, timezone
import openpyxl

ROOT=pathlib.Path(__file__).resolve().parents[2]
BIO=ROOT/"BIO_CHI"
SRC=BIO/"config"/"SU2026_SUPPDATA_WORKBOOK_FREEZE_v0_1.json"
CFG=BIO/"config"/"SU2026_CROSS_CELL_CARRIER_COVERAGE_FREEZE_v0_1.json"
OUTDIR=BIO/"artifacts"/"generated"
OUTDIR.mkdir(parents=True,exist_ok=True)
OUT=OUTDIR/"su2026_cross_cell_carrier_coverage_v0_1.json"
UA="BioChiReviewerReproducibility/0.1"

def fetch(url):
    req=urllib.request.Request(url,headers={"User-Agent":UA,"Accept":"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,application/octet-stream,*/*;q=0.5"})
    with urllib.request.urlopen(req,timeout=180) as r: return r.read()
def md5(b): return hashlib.md5(b).hexdigest()
def sha256(b): return hashlib.sha256(b).hexdigest()
def acquire(cfg):
    errs=[]
    for u in cfg["candidate_urls"]:
        try:
            b=fetch(u)
            if md5(b)==cfg["expected_md5"]: return b
            errs.append({"url":u,"md5":md5(b)})
        except Exception as e: errs.append({"url":u,"error":f"{type(e).__name__}: {e}"})
    raise RuntimeError(json.dumps(errs))
def num(v,coord):
    if isinstance(v,bool) or not isinstance(v,(int,float)): raise ValueError(f"{coord} not numeric")
    x=float(v)
    if not math.isfinite(x): raise ValueError(f"{coord} nonfinite")
    return x

def contribution_map(ws):
    out={}
    dup=[]
    for r in range(3,ws.max_row+1):
        g=ws.cell(r,1).value
        if g is None: continue
        g=str(g).strip()
        if not g: continue
        if g in out: dup.append(g); continue
        out[g]=[num(ws.cell(r,2).value,ws.cell(r,2).coordinate),num(ws.cell(r,3).value,ws.cell(r,3).coordinate)]
    if dup: raise RuntimeError("duplicate contribution genes: "+repr(sorted(set(dup))[:20]))
    return out

def select(mp,idx,pol):
    vals=[]
    for g,scores in mp.items():
        s=scores[idx]
        if pol=="positive" and s>0: vals.append((g,s))
        if pol=="negative" and s<0: vals.append((g,s))
    vals.sort(key=(lambda x:(-x[1],x[0])) if pol=="positive" else (lambda x:(x[1],x[0])))
    genes=[g for g,_ in vals[:500]]
    if len(genes)!=500: raise RuntimeError(f"{idx} {pol} selected {len(genes)}")
    return genes

def gene_ids(ws):
    genes=[]; seen=set(); dup=[]
    for r in range(3,ws.max_row+1):
        v=ws.cell(r,1).value
        if v is None: continue
        g=str(v).strip()
        if not g: continue
        if g in seen: dup.append(g)
        else: seen.add(g); genes.append(g)
    return genes,sorted(set(dup))

def h(xs): return hashlib.sha256("\n".join(xs).encode()).hexdigest()

def main():
    src=json.loads(SRC.read_text())
    cfg=json.loads(CFG.read_text())
    data=acquire(src)
    if sha256(data)!=cfg["source_workbook_sha256"]: raise SystemExit("workbook SHA mismatch")
    path=OUTDIR/src["expected_filename"]; path.write_bytes(data)
    wb=openpyxl.load_workbook(path,read_only=False,data_only=False)
    cm=contribution_map(wb["Supplementary Data 4"])
    sets={
      "Mearly_positive":select(cm,1,"positive"),
      "Mearly_negative":select(cm,1,"negative"),
      "Mlate_positive":select(cm,0,"positive"),
      "Mlate_negative":select(cm,0,"negative")
    }
    systems={}
    for srcdef in cfg["candidate_transport_sources"]:
        ws=wb[srcdef["sheet"]]
        ids,dup=gene_ids(ws); universe=set(ids)
        cov={}
        for name,genes in sets.items():
            present=[g for g in genes if g in universe]
            missing=[g for g in genes if g not in universe]
            cov[name]={
              "source_n":len(genes),"present_n":len(present),"missing_n":len(missing),
              "coverage_fraction":len(present)/len(genes),
              "source_genes_sha256":h(genes),
              "present_genes_sha256":h(present),
              "missing_genes":missing
            }
        systems[srcdef["system"]]={
          "sheet":srcdef["sheet"],"treatment":srcdef["treatment"],
          "gene_universe_n":len(ids),"duplicate_gene_ids":dup,
          "coverage":cov
        }
    wb.close()
    result={
      "schema_version":"0.1","generated_at_utc":datetime.now(timezone.utc).isoformat(),
      "gate":cfg["gate"],"status":"PASS_COVERAGE_AUDIT",
      "workbook_sha256":sha256(data),
      "expression_values_read":False,
      "source_gene_sets":{k:{"n":len(v),"sha256":h(v)} for k,v in sets.items()},
      "systems":systems
    }
    OUT.write_text(json.dumps(result,indent=2)+"\n")
    print(json.dumps({
      "status":result["status"],
      "systems":{s:{"gene_universe_n":v["gene_universe_n"],
                     "duplicates":len(v["duplicate_gene_ids"]),
                     "coverage":{k:{"present_n":c["present_n"],"coverage_fraction":c["coverage_fraction"]} for k,c in v["coverage"].items()}}
                 for s,v in systems.items()},
      "expression_values_read":False
    },indent=2))
    print("BIO_CHI_SU2026_CROSS_CELL_COVERAGE_PASS")
if __name__=="__main__": main()
