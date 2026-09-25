#!/usr/bin/env python3
from __future__ import annotations

import hashlib
import json
import pathlib
import urllib.request
from datetime import datetime, timezone

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parents[2]
BIO = ROOT / "BIO_CHI"
CFG = BIO / "config" / "SU2026_SUPPDATA_WORKBOOK_FREEZE_v0_1.json"
OUTDIR = BIO / "artifacts" / "generated"
OUTDIR.mkdir(parents=True, exist_ok=True)
OUT = OUTDIR / "su2026_suppdata_workbook_freeze_v0_1.json"

UA = "BioChiReviewerReproducibility/0.1"

def fetch(url: str) -> tuple[bytes, dict]:
    req = urllib.request.Request(
        url,
        headers={
            "User-Agent": UA,
            "Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,application/octet-stream,*/*;q=0.5",
        },
    )
    with urllib.request.urlopen(req, timeout=180) as r:
        data = r.read()
        return data, {
            "requested_url": url,
            "resolved_url": r.geturl(),
            "content_type": r.headers.get("Content-Type"),
            "content_length": r.headers.get("Content-Length"),
            "etag": r.headers.get("ETag"),
            "last_modified": r.headers.get("Last-Modified"),
        }

def md5(data: bytes) -> str:
    return hashlib.md5(data).hexdigest()

def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()

def classify_cell(cell):
    v = cell.value
    if v is None:
        return None
    if isinstance(v, str):
        # Preserve text only. Formula strings are not recorded.
        if v.startswith("="):
            return {"coord": cell.coordinate, "kind": "FORMULA"}
        s = v.strip()
        if not s:
            return None
        return {"coord": cell.coordinate, "kind": "TEXT", "text": s}
    if cell.is_date:
        return {"coord": cell.coordinate, "kind": "DATE"}
    if isinstance(v, bool):
        return {"coord": cell.coordinate, "kind": "BOOLEAN"}
    if isinstance(v, (int, float)):
        return {"coord": cell.coordinate, "kind": "NUMERIC"}
    return {"coord": cell.coordinate, "kind": type(v).__name__.upper()}

def main():
    cfg = json.loads(CFG.read_text(encoding="utf-8"))
    data = None
    meta = None
    errors = []
    for url in cfg["candidate_urls"]:
        try:
            candidate, candidate_meta = fetch(url)
            if md5(candidate) != cfg["expected_md5"]:
                errors.append({
                    "url": url,
                    "error": "MD5_MISMATCH",
                    "observed_md5": md5(candidate),
                    "bytes": len(candidate),
                })
                continue
            data = candidate
            meta = candidate_meta
            break
        except Exception as exc:
            errors.append({"url": url, "error": f"{type(exc).__name__}: {exc}"})
    if data is None:
        raise SystemExit("SU2026_SUPPDATA_DOWNLOAD_OR_MD5_FAILURE: " + json.dumps(errors))

    work = OUTDIR / cfg["expected_filename"]
    work.write_bytes(data)

    wb = openpyxl.load_workbook(work, read_only=True, data_only=False)
    sheets = []
    for ws in wb.worksheets:
        cells = []
        for row in ws.iter_rows():
            for cell in row:
                item = classify_cell(cell)
                if item is not None:
                    cells.append(item)
        sheets.append({
            "title": ws.title,
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "nonempty_cell_count": len(cells),
            "schema_cells": cells,
        })
    wb.close()

    result = {
        "schema_version": "0.1",
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "gate": cfg["gate"],
        "status": "PASS_SOURCE_BYTES_AND_TEXT_ONLY_SCHEMA_INVENTORY",
        "filename": cfg["expected_filename"],
        "byte_size": len(data),
        "md5": md5(data),
        "sha256": sha256(data),
        "download_meta": meta,
        "download_attempts": errors,
        "sheet_count": len(sheets),
        "sheets": sheets,
        "numeric_values_recorded": False,
        "formulas_recorded": False,
        "scientific_values_interpreted": False,
        "next_gate": "freeze exact sheet/range and source-defined scenario-to-parameter semantic mapping from text schema before opening numeric coefficient values",
    }
    OUT.write_text(json.dumps(result, indent=2) + "\n", encoding="utf-8")

    print(json.dumps({
        "status": result["status"],
        "filename": result["filename"],
        "byte_size": result["byte_size"],
        "md5": result["md5"],
        "sha256": result["sha256"],
        "sheet_count": result["sheet_count"],
        "sheets": [
            {
                "title": s["title"],
                "max_row": s["max_row"],
                "max_column": s["max_column"],
                "text_cells": [c for c in s["schema_cells"] if c["kind"] == "TEXT"][:80],
                "nonempty_cell_count": s["nonempty_cell_count"],
            }
            for s in sheets
        ],
        "numeric_values_recorded": False,
    }, indent=2))
    print("BIO_CHI_SU2026_SUPPDATA_WORKBOOK_FREEZE_PASS")

if __name__ == "__main__":
    main()
