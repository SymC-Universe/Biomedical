#!/usr/bin/env python3
from __future__ import annotations
import hashlib, json, math, pathlib, statistics, urllib.request
from collections import defaultdict
from datetime import datetime, timezone

import openpyxl

ROOT=pathlib.Path(__file__).resolve().parents[2]
BIO=ROOT/"BIO_CHI"
CFG=BIO/"config"/"HARMANGE2023_SOURCE_DATA_WORKBOOK_FREEZE_v0_1.json"
EVENT=BIO/"config"/"HARMANGE2023_TOPDOWN_EVENT_FREEZE_v0_1.json"
OUTDIR=BIO/"artifacts"/"generated"
OUTDIR.mkdir(parents=True,exist_ok=True)
OUT=OUTDIR/"harmange2023_topdown_analysis_v0_1.json"
AUDIT=OUTDIR/"HARMANGE2023_TOPDOWN_ANALYSIS_V01_AUDIT.md"
UA="BioChiReviewerReproducibility/0.1"

def fetch(url):
    req=urllib.request.Request(url,headers={"User-Agent":UA,"Accept":"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,*/*;q=0.5"})
    with urllib.request.urlopen(req,timeout=180) as r:
        return r.read()

def md5(b): return hashlib.md5(b).hexdigest()
def sha256(b): return hashlib.sha256(b).hexdigest()

def num(v,coord,allow_zero=True):
    if isinstance(v,bool) or not isinstance(v,(int,float)):
        raise ValueError(f"{coord}: expected numeric, got {v!r}")
    x=float(v)
    if not math.isfinite(x): raise ValueError(f"{coord}: nonfinite")
    if not allow_zero and x==0: raise ValueError(f"{coord}: zero not allowed")
    return x

def exact_one_sided_sign_p(successes, failures):
    n=successes+failures
    if n==0: return None
    return sum(math.comb(n,k) for k in range(successes,n+1))/(2**n)

def lineage_pairs(ws,treatment,direction):
    headers=[ws.cell(1,c).value for c in range(1,5)]
    if headers != ["Condition","Lineage","propOfprimed","numCells"]:
        raise RuntimeError(f"{ws.title}: header mismatch {headers!r}")
    d=defaultdict(dict)
    for r in range(2,ws.max_row+1):
        cond=ws.cell(r,1).value
        lin=ws.cell(r,2).value
        if cond is None or lin is None: continue
        cond=str(cond).strip(); lin=str(lin).strip()
        prop=num(ws.cell(r,3).value,ws.cell(r,3).coordinate)
        nc=num(ws.cell(r,4).value,ws.cell(r,4).coordinate)
        if not (0<=prop<=1): raise RuntimeError(f"{ws.title} {lin} {cond}: prop outside [0,1]")
        if nc<0: raise RuntimeError(f"{ws.title} {lin} {cond}: negative numCells")
        if cond in d[lin]: raise RuntimeError(f"{ws.title}: duplicate lineage-condition {lin}/{cond}")
        d[lin][cond]={"prop":prop,"numCells":nc}
    common=sorted(lin for lin,x in d.items() if "Control" in x and treatment in x)
    rows=[]
    for lin in common:
        c=d[lin]["Control"]["prop"]; t=d[lin][treatment]["prop"]; delta=t-c
        rows.append({"lineage":lin,"control":c,"treatment":t,"delta":delta,
                     "control_numCells":d[lin]["Control"]["numCells"],"treatment_numCells":d[lin][treatment]["numCells"]})
    pos=sum(x["delta"]>0 for x in rows); neg=sum(x["delta"]<0 for x in rows); zero=sum(x["delta"]==0 for x in rows)
    successes=neg if direction=="negative" else pos
    failures=pos if direction=="negative" else neg
    return {
      "treatment":treatment,"direction":direction,"paired_lineage_count":len(rows),
      "median_delta":statistics.median([x["delta"] for x in rows]) if rows else None,
      "positive_delta_count":pos,"negative_delta_count":neg,"zero_delta_count":zero,
      "one_sided_exact_sign_p":exact_one_sided_sign_p(successes,failures),
      "lineages":rows
    }

def replicate_means(ws,condition_col,value_col,rep_col):
    d=defaultdict(lambda:defaultdict(list))
    for r in range(2,ws.max_row+1):
        cond=ws.cell(r,condition_col).value
        rep=ws.cell(r,rep_col).value
        val=ws.cell(r,value_col).value
        if cond is None or rep is None or val is None: continue
        cond=str(cond).strip(); rep=str(rep).strip()
        x=num(val,ws.cell(r,value_col).coordinate)
        if x<0: raise RuntimeError(f"{ws.title}: negative outcome")
        d[cond][rep].append(x)
    out={}
    for cond,reps in d.items():
        out[cond]={rep:{"mean":sum(vals)/len(vals),"n_technical":len(vals),"values":vals} for rep,vals in reps.items()}
    return out

def compare_replicates(groups,treatment="PI3Ki",control="Control"):
    if treatment not in groups or control not in groups:
        raise RuntimeError(f"missing treatment/control: {groups.keys()}")
    common=sorted(set(groups[treatment]) & set(groups[control]))
    rows=[]
    for rep in common:
        t=groups[treatment][rep]["mean"]; c=groups[control][rep]["mean"]; delta=t-c
        rows.append({"replicate":rep,"control_mean":c,"treatment_mean":t,"delta":delta,
                     "control_n_technical":groups[control][rep]["n_technical"],"treatment_n_technical":groups[treatment][rep]["n_technical"]})
    neg=sum(x["delta"]<0 for x in rows); pos=sum(x["delta"]>0 for x in rows); zero=sum(x["delta"]==0 for x in rows)
    return {
      "matched_replicate_count":len(rows),"replicates":rows,
      "negative_delta_count":neg,"positive_delta_count":pos,"zero_delta_count":zero,
      "all_treatment_lower":bool(rows) and neg==len(rows),
      "median_delta":statistics.median([x["delta"] for x in rows]) if rows else None,
      "one_sided_exact_sign_p":exact_one_sided_sign_p(neg,pos)
    }

def main():
    cfg=json.loads(CFG.read_text())
    event=json.loads(EVENT.read_text())
    data=fetch(cfg["source"]["url"])
    if md5(data)!=cfg["source"]["reported_md5"]: raise SystemExit("workbook MD5 mismatch")
    if sha256(data)!=event["source_workbook_sha256"]: raise SystemExit("workbook SHA256 mismatch")
    path=OUTDIR/cfg["source"]["filename"]; path.write_bytes(data)
    wb=openpyxl.load_workbook(path,read_only=False,data_only=False)

    switch_out=lineage_pairs(wb["Fig. 5H"],"PI3Ki","negative")
    switch_in=lineage_pairs(wb["Fig. 5G"],"TGFB1","positive")
    # Frozen primary rule explicitly expects at least 6/8 for Fig 5H.
    switch_out_pass=(switch_out["paired_lineage_count"]==8 and switch_out["median_delta"]<0 and switch_out["negative_delta_count"]>=6)

    ws6c=wb["Fig. 6C"]
    h6c=[ws6c.cell(1,c).value for c in range(1,4)]
    if h6c!=["Condition","replicate","ncells"]: raise RuntimeError(f"Fig.6C header mismatch {h6c!r}")
    cells_groups=replicate_means(ws6c,1,3,2)
    cells_cmp=compare_replicates(cells_groups)

    ws6d=wb["Fig. 6D"]
    h6d=[ws6d.cell(1,c).value for c in range(1,5)]
    if h6d!=["condition","replicate","well","numColonies"]: raise RuntimeError(f"Fig.6D header mismatch {h6d!r}")
    colonies_groups=replicate_means(ws6d,1,4,2)
    colonies_cmp=compare_replicates(colonies_groups)

    functional_pass=cells_cmp["all_treatment_lower"] and colonies_cmp["all_treatment_lower"]
    biochi_pass=switch_out_pass and functional_pass

    result={
      "schema_version":"0.1","generated_at_utc":datetime.now(timezone.utc).isoformat(),
      "gate":"Harmange 2023 top-down state-switch-to-resistance event",
      "status":"PASS_ANALYSIS_EXECUTED",
      "source_workbook":{"sha256":sha256(data),"md5":md5(data),"bytes":len(data)},
      "state_switch_primary":{"result":switch_out,"gate_pass":switch_out_pass},
      "reverse_switch_context":switch_in,
      "functional_outcome":{
        "resistant_cells":{"comparison":cells_cmp,"all_groups":cells_groups},
        "resistant_colonies":{"comparison":colonies_cmp,"all_groups":colonies_groups},
        "gate_pass":functional_pass
      },
      "bio_chi":{
        "gate_pass":biochi_pass,
        "disposition":"P0Q_WITHIN_SOURCE_STATE_MEMORY_TO_RESISTANCE_RELATION_SUPPORTED" if biochi_pass else "P0Q_STATE_MEMORY_TO_RESISTANCE_RELATION_NOT_REPRODUCED_UNDER_FROZEN_GATE"
      },
      "downward_representation":{
        "minimum_internal_representation":"lineage_resolved_two_state_composition",
        "full_transcriptomic_modal_analysis_opened":False,
        "scalar_opened":False,
        "scalar_disposition":"SCALAR_NOT_REQUIRED_AND_NOT_LICENSED"
      },
      "epistemic":{
        "tier":"P0-Q","independent_P1_confirmation":False,
        "lineages_are_not_biological_replicates":True,
        "biological_replicate_means_used_for_resistance_direction_gate":True,
        "sign_p_values_are_resolution_consistency_diagnostics":True
      }
    }
    wb.close()
    OUT.write_text(json.dumps(result,indent=2)+"\n")

    lines=[
      "# Harmange 2023 top-down analysis v0.1","",
      f"**Bio Chi disposition:** {result['bio_chi']['disposition']}","",
      "## Lineage state switching","",
      f"- PI3Ki paired lineages: {switch_out['paired_lineage_count']}",
      f"- PI3Ki median change in primed fraction: {switch_out['median_delta']:.8g}",
      f"- PI3Ki negative/positive/zero lineages: {switch_out['negative_delta_count']}/{switch_out['positive_delta_count']}/{switch_out['zero_delta_count']}",
      f"- PI3Ki frozen lineage gate: {'PASS' if switch_out_pass else 'FAIL'}",
      f"- TGFB1 reverse-direction paired lineages: {switch_in['paired_lineage_count']}",
      f"- TGFB1 median change in primed fraction: {switch_in['median_delta']:.8g}","",
      "## Functional resistance","",
      f"- Resistant-cell matched biological replicates: {cells_cmp['matched_replicate_count']}",
      f"- Resistant-cell PI3Ki lower in all replicates: {cells_cmp['all_treatment_lower']}",
      f"- Resistant-colony matched biological replicates: {colonies_cmp['matched_replicate_count']}",
      f"- Resistant-colony PI3Ki lower in all replicates: {colonies_cmp['all_treatment_lower']}",
      f"- Frozen functional gate: {'PASS' if functional_pass else 'FAIL'}","",
      "## Downward representation","",
      "The whole-event analysis requires lineage-resolved two-state composition, not an invented scalar. Full transcriptomic PCA is not opened merely to create an additional representation."
    ]
    AUDIT.write_text("\n".join(lines)+"\n")
    print(json.dumps({
      "status":result["status"],"state_switch_primary":{"paired":switch_out["paired_lineage_count"],"median_delta":switch_out["median_delta"],"neg":switch_out["negative_delta_count"],"pos":switch_out["positive_delta_count"],"zero":switch_out["zero_delta_count"],"p":switch_out["one_sided_exact_sign_p"],"gate_pass":switch_out_pass},
      "reverse_context":{"paired":switch_in["paired_lineage_count"],"median_delta":switch_in["median_delta"],"pos":switch_in["positive_delta_count"],"neg":switch_in["negative_delta_count"],"zero":switch_in["zero_delta_count"],"p":switch_in["one_sided_exact_sign_p"]},
      "functional":{"cells":cells_cmp,"colonies":colonies_cmp,"gate_pass":functional_pass},
      "bio_chi":result["bio_chi"],"downward_representation":result["downward_representation"]
    },indent=2))
    print("BIO_CHI_HARMANGE_TOPDOWN_ANALYSIS_DONE")

if __name__=="__main__": main()
