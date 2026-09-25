#!/usr/bin/env python3
from __future__ import annotations
import hashlib,json,pathlib,urllib.request
from datetime import datetime,timezone
import openpyxl

ROOT=pathlib.Path(__file__).resolve().parents[2]
BIO=ROOT/"BIO_CHI"; CFG=BIO/"config"/"HARMANGE2023_SOURCE_DATA_WORKBOOK_FREEZE_v0_1.json"
OUTDIR=BIO/"artifacts"/"generated"; OUTDIR.mkdir(parents=True,exist_ok=True)
OUT=OUTDIR/"harmange2023_source_data_workbook_freeze_v0_1.json"
UA="BioChiReviewerReproducibility/0.1"

def fetch(url):
 req=urllib.request.Request(url,headers={"User-Agent":UA,"Accept":"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,application/octet-stream,*/*;q=0.5"})
 with urllib.request.urlopen(req,timeout=180) as r:
  b=r.read()
  return b,{"resolved_url":r.geturl(),"content_type":r.headers.get("Content-Type"),"content_length":r.headers.get("Content-Length"),"last_modified":r.headers.get("Last-Modified")}

def md5(b):return hashlib.md5(b).hexdigest()
def sha(b):return hashlib.sha256(b).hexdigest()

def cell_schema(cell):
 v=cell.value
 if v is None:return None
 if isinstance(v,str):
  if v.startswith("="):return {"coord":cell.coordinate,"kind":"FORMULA"}
  s=v.strip()
  return None if not s else {"coord":cell.coordinate,"kind":"TEXT","text":s}
 if cell.is_date:return {"coord":cell.coordinate,"kind":"DATE"}
 if isinstance(v,bool):return {"coord":cell.coordinate,"kind":"BOOLEAN"}
 if isinstance(v,(int,float)):return {"coord":cell.coordinate,"kind":"NUMERIC"}
 return {"coord":cell.coordinate,"kind":type(v).__name__.upper()}

def main():
 cfg=json.loads(CFG.read_text())
 data,meta=fetch(cfg["source"]["url"])
 if md5(data)!=cfg["source"]["reported_md5"]:
  raise SystemExit(f"MD5 mismatch {md5(data)}")
 path=OUTDIR/cfg["source"]["filename"];path.write_bytes(data)
 wb=openpyxl.load_workbook(path,read_only=True,data_only=False)
 sheets=[]
 for ws in wb.worksheets:
  cells=[]
  for row in ws.iter_rows():
   for cell in row:
    x=cell_schema(cell)
    if x:cells.append(x)
  sheets.append({
   "title":ws.title,"max_row":ws.max_row,"max_column":ws.max_column,
   "nonempty_cell_count":len(cells),"schema_cells":cells
  })
 wb.close()
 result={
  "schema_version":"0.1","generated_at_utc":datetime.now(timezone.utc).isoformat(),
  "status":"PASS_SOURCE_BYTES_AND_TEXT_ONLY_SCHEMA_INVENTORY",
  "filename":cfg["source"]["filename"],"byte_size":len(data),"md5":md5(data),"sha256":sha(data),
  "download_meta":meta,"sheet_count":len(sheets),"sheets":sheets,
  "numeric_values_recorded":False,"scientific_values_interpreted":False
 }
 OUT.write_text(json.dumps(result,indent=2)+"\n")
 print(json.dumps({
  "status":result["status"],"byte_size":len(data),"md5":result["md5"],"sha256":result["sha256"],
  "sheet_count":len(sheets),
  "sheets":[{"title":s["title"],"max_row":s["max_row"],"max_column":s["max_column"],
             "text_cells":[c for c in s["schema_cells"] if c["kind"]=="TEXT"][:120],
             "nonempty_cell_count":s["nonempty_cell_count"]} for s in sheets],
  "numeric_values_recorded":False
 },indent=2))
 print("BIO_CHI_HARMANGE_SOURCE_DATA_WORKBOOK_FREEZE_PASS")
if __name__=="__main__":main()
